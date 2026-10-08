from datetime import date
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from apps.core.testing import make_animal, make_person, next_rut
from apps.core.testing_api import PASSWORD, ApiClient, signed_in
from apps.operations.models import BackupRun
from apps.organization.models import Locality
from apps.people.models import Membership, ReviewStatus, TeamMember, User


class SessionTests(TestCase):
    def test_login_me_logout(self):
        _, user = signed_in("worker", email="diego@example.org")
        client = ApiClient()
        self.assertEqual(client.post_json("/api/auth/login", {"email": "diego@example.org", "password": "mala"}).status_code, 400)
        me = client.post_json("/api/auth/login", {"email": "DIEGO@example.org", "password": PASSWORD}).json()
        self.assertEqual((me["email"], me["role"]), ("diego@example.org", "worker"))
        self.assertEqual(client.get("/api/auth/me").status_code, 200)
        client.post_json("/api/auth/logout")
        self.assertEqual(client.get("/api/auth/me").status_code, 401)

    def test_changes_require_the_csrf_token(self):
        _, user = signed_in("worker", email="csrf@example.org")
        client = ApiClient(enforce_csrf_checks=True)
        client.login(username="csrf@example.org", password=PASSWORD)
        url = "/api/dashboard/traits"
        self.assertEqual(client.post_json(url, {"label": "Sociable"}).status_code, 403)
        token = client.get("/api/auth/csrf").json()["csrf_token"]
        response = client.post(url, '{"label": "Sociable"}', content_type="application/json", HTTP_X_CSRFTOKEN=token)
        self.assertEqual(response.status_code, 201)


class ApplicationsTests(TestCase):
    def setUp(self):
        self.public = ApiClient()

    def apply(self, **extra):
        data = {
            "rut": next_rut(), "first_names": "Sofía", "paternal_surname": "Navarro", "phone": "+56 9 7654 3210",
            "locality_id": Locality.objects.first().id, "consent": True, **extra,
        }
        return self.public.post_json("/api/public/applications", data)

    def test_member_and_volunteer_start_pending(self):
        response = self.apply(membership={"kind": "donation", "monthly_amount": 10000}, volunteer={"availability_ids": []})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Membership.objects.get().status, ReviewStatus.PENDING)
        self.assertEqual(TeamMember.objects.get().status, ReviewStatus.PENDING)

    def test_sponsorship_needs_a_visible_animal(self):
        self.assertEqual(self.apply(membership={"kind": "sponsorship", "monthly_amount": 7000}).status_code, 400)
        animal = make_animal()
        ok = self.apply(membership={"kind": "sponsorship", "monthly_amount": 7000, "animal_id": animal.id})
        self.assertEqual(ok.status_code, 201)

    def test_only_administration_approves(self):
        self.apply(volunteer={"availability_ids": []})
        member = TeamMember.objects.get()
        worker, _ = signed_in("worker")
        self.assertEqual(worker.post_json(f"/api/dashboard/team-members/{member.id}/approve").status_code, 403)
        admin, _ = signed_in("admin")
        approved = admin.post_json(f"/api/dashboard/team-members/{member.id}/approve")
        self.assertEqual(approved.status_code, 200)
        member.refresh_from_db()
        self.assertEqual(member.joined_at, date.today())


class TeamExitTests(TestCase):
    def setUp(self):
        self.admin, self.admin_user = signed_in("admin")
        self.volunteer, self.volunteer_user = signed_in("worker")
        self.member = TeamMember.objects.create(
            person=self.volunteer_user.person, status=ReviewStatus.APPROVED, joined_at=date(2026, 1, 1), is_foster_home=True
        )

    def test_cannot_leave_while_fostering_and_leaving_revokes_access(self):
        animal = make_animal(name="Canela", foster_home=self.member)
        url = f"/api/dashboard/team-members/{self.member.id}/leave"
        blocked = self.admin.post_json(url, {"effective_date": str(date.today())})
        self.assertIn("Canela", blocked.json()["detail"])
        self.admin.put_json(f"/api/dashboard/animals/{animal.id}/foster-home", {"team_member_id": None})
        self.assertEqual(self.admin.post_json(url, {"effective_date": str(date.today())}).status_code, 200)
        self.volunteer_user.refresh_from_db()
        self.assertFalse(self.volunteer_user.is_active)
        self.assertEqual(self.volunteer.get("/api/dashboard/summary").status_code, 401)

    def test_last_administration_account_cannot_be_revoked(self):
        other_admin = User.objects.create_user(make_person(email="otra@example.org"), PASSWORD, role="admin")
        response = self.admin.post_json(f"/api/dashboard/people/{other_admin.person_id}/revoke-access")
        self.assertEqual(response.status_code, 200)
        self_revoke = self.admin.post_json(f"/api/dashboard/people/{self.admin_user.person_id}/revoke-access")
        self.assertEqual(self_revoke.status_code, 400)

    def test_rut_cannot_be_edited(self):
        person = self.volunteer_user.person
        self.admin.patch_json(f"/api/dashboard/people/{person.id}", {"document_number": "11111111-1", "phone": "+56 9 1111 2222"})
        person.refresh_from_db()
        self.assertEqual(person.phone, "+56 9 1111 2222")
        self.assertNotEqual(person.document_number, "11111111-1")


class RestrictionsAndBackupsTests(TestCase):
    def test_team_checks_admin_adds(self):
        worker, _ = signed_in("worker")
        admin, _ = signed_in("admin")
        rut_value = next_rut()
        data = {"rut": rut_value, "first_names": "Ramón", "paternal_surname": "Ferrer", "reason": "abuse", "notes": "Denuncia 2024"}
        self.assertEqual(worker.post_json("/api/dashboard/restrictions", data).status_code, 403)
        self.assertEqual(admin.post_json("/api/dashboard/restrictions", data).status_code, 201)
        check = worker.get("/api/dashboard/restrictions/check", {"rut": rut_value}).json()
        self.assertEqual((check["can_adopt"], check["name"]), (False, "Ramón Ferrer"))
        invalid = worker.get("/api/dashboard/restrictions/check", {"rut": "12.345.678-0"})
        self.assertEqual(invalid.status_code, 400)

    def test_excel_backup(self):
        admin, _ = signed_in("admin")
        make_animal(name="Bruno")
        response = admin.get("/api/dashboard/backups/excel")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(book.sheetnames, ["Animales", "Adopciones", "Personas", "Socios", "Equipo", "Restricciones", "Avisos"])
        self.assertEqual(BackupRun.objects.get().format, BackupRun.Format.EXCEL)
