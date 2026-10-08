"""Excel imports: template, validation of every row, and an all-or-nothing save.

Each kind declares its columns. A file is accepted only if its header matches the template exactly and every row is valid;
otherwise nothing is saved and the preview explains every problem.
"""

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Callable

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.adoptions.models import Adoption
from apps.animals.models import Animal, Species
from apps.animals.services import in_care, unique_slug
from apps.core import rut
from apps.people.models import Membership, Person, ReviewStatus, TeamMember
from apps.restrictions.models import Restriction

from .models import ImportBatch


class CellError(Exception):
    pass


# ---------- Cell parsers ----------

def text(value, required=True):
    value = "" if value is None else str(value).strip()
    if required and not value:
        raise CellError("Falta este dato.")
    return value


def optional_text(value):
    return text(value, required=False)


def phone_cell(value):
    value = text(value)
    if sum(ch.isdigit() for ch in value) < 9:
        raise CellError("Escribe un teléfono de al menos 9 números.")
    return value


def email_cell(value):
    from django.core.validators import validate_email

    value = optional_text(value).lower()
    if value:
        try:
            validate_email(value)
        except ValidationError:
            raise CellError("Correo no válido.") from None
    return value


def rut_cell(value):
    try:
        return rut.normalize(text(value))
    except ValidationError as exc:
        raise CellError(f"RUT no válido. {exc.message}") from None


def date_cell(value, required=True):
    if value in (None, ""):
        if required:
            raise CellError("Falta la fecha.")
        return None
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    else:
        for fmt in ("%d-%m-%Y", "%d/%m/%Y"):
            try:
                result = datetime.strptime(str(value).strip(), fmt).date()
                break
            except ValueError:
                continue
        else:
            raise CellError("Fecha no válida. Escríbela como día-mes-año, por ejemplo 15-03-2024.")
    if result > date.today():
        raise CellError("La fecha no puede ser futura.")
    return result


def optional_date(value):
    return date_cell(value, required=False)


def amount_cell(value, required=True):
    if value in (None, ""):
        if required:
            raise CellError("Falta el monto.")
        return None
    try:
        number = float(str(value).replace("$", "").replace(".", "").replace(",", ".").strip()) if isinstance(value, str) else float(value)
    except ValueError:
        raise CellError("Escribe el monto solo con números, por ejemplo 30000.") from None
    if number < 0 or number != int(number):
        raise CellError("El monto debe ser un número entero mayor o igual a 0.")
    return int(number)


def optional_amount(value):
    return amount_cell(value, required=False)


def choice(options):
    lookup = {k.lower(): v for k, v in options.items()}

    def parse(value):
        key = text(value).lower()
        if key not in lookup:
            raise CellError(f"Valor no permitido. Usa: {', '.join(options)}.")
        return lookup[key]

    return parse


def species_cell(value, allow_all=False):
    name = text(value)
    if allow_all and name.lower() in {"todas", "todas las especies"}:
        return None
    species = Species.objects.filter(name__iexact=name).first() or Species.objects.filter(plural_name__iexact=name).first()
    if not species:
        raise CellError(f"La especie “{name}” no existe. Créala en Apariencia › Especies o corrige la fila.")
    return species


SEX = {"Hembra": Animal.Sex.FEMALE, "Macho": Animal.Sex.MALE, "Desconocido": Animal.Sex.UNKNOWN}
SIZE = {"Pequeño": Animal.Size.SMALL, "Mediano": Animal.Size.MEDIUM, "Grande": Animal.Size.LARGE}
PERSON_TYPE = {"Socio": "donation", "Padrino": "sponsorship", "Voluntario": "volunteer", "Planta": "staff"}
REASON = {label: value for value, label in Restriction.Reason.choices}


@dataclass
class Column:
    header: str
    key: str
    parse: Callable
    required: bool = True
    hint: str = ""


PERSON_COLUMNS = [
    Column("RUT", "rut", rut_cell, hint="Con dígito verificador, por ejemplo 12.345.678-5"),
    Column("Nombres", "first_names", text),
    Column("Primer apellido", "paternal_surname", text),
    Column("Segundo apellido", "maternal_surname", optional_text, required=False),
]

ANIMAL_DATES = [
    Column("Fecha de nacimiento aproximada", "birth_date_approx", date_cell, hint="Día-mes-año"),
    Column("Fecha de llegada", "intake_date", date_cell, hint="Día-mes-año"),
]


@dataclass
class Kind:
    label: str
    columns: list[Column]
    row_check: Callable = None
    file_check: Callable = None
    save: Callable = None
    notes: list[str] = field(default_factory=list)


# ---------- Row and file rules ----------

def _check_animal_dates(row):
    if row["intake_date"] < row["birth_date_approx"]:
        raise CellError("La fecha de llegada no puede ser anterior a la de nacimiento.")


def _check_adoption_row(row):
    _check_animal_dates(row)
    if row["contract_signed_at"] < row["intake_date"]:
        raise CellError("El contrato no puede ser anterior a la llegada del animal.")
    if row["returned_at"] and row["returned_at"] < row["contract_signed_at"]:
        raise CellError("La devolución no puede ser anterior al contrato.")
    if row["returned_at"] and not row["return_reason"]:
        raise CellError("Si hubo devolución, escribe el motivo.")
    if row["return_reason"] and not row["returned_at"]:
        raise CellError("Hay un motivo de devolución pero falta la fecha de devolución.")


def _animal_key(row):
    return (row["animal_name"].lower(), row["species"].id, row["birth_date_approx"], row["intake_date"])


def _check_adoption_file(rows):
    """Rows for the same animal (same name, species, birth and arrival) form its history, oldest first."""
    errors = defaultdict(list)
    groups = defaultdict(list)
    for number, row in rows:
        groups[_animal_key(row)].append((number, row))
    for history in groups.values():
        history.sort(key=lambda item: item[1]["contract_signed_at"])
        for (n1, r1), (n2, r2) in zip(history, history[1:]):
            if not r1["returned_at"]:
                errors[n1].append(f"{r1['animal_name']} tiene otra adopción posterior (fila {n2}): esta debe tener fecha de devolución.")
            elif r2["contract_signed_at"] < r1["returned_at"]:
                errors[n2].append(f"Este contrato es anterior a la devolución de la fila {n1}.")
    return errors


def _check_person_row(row):
    if row["email"] and Person.objects.filter(email=row["email"]).exclude(document_number=row["rut"]).exists():
        raise CellError("Ese correo ya pertenece a otra persona.")
    kind = row["type"]
    if kind in {"donation", "sponsorship"} and not row["monthly_amount"]:
        raise CellError("Un socio o padrino necesita un monto mensual mayor a 0.")
    if kind == "sponsorship":
        if not row["sponsored_animal"]:
            raise CellError("Un padrino necesita el nombre del animal que apadrina.")
        matches = list(in_care().filter(name__iexact=row["sponsored_animal"])[:2])
        if not matches:
            raise CellError(f"No hay un animal a cargo de la fundación llamado “{row['sponsored_animal']}”.")
        if len(matches) > 1:
            raise CellError(f"Hay más de un animal llamado “{row['sponsored_animal']}”. Cambia el nombre de uno en el panel e intenta de nuevo.")
        row["animal"] = matches[0]
    elif row["sponsored_animal"]:
        raise CellError("Solo un padrino lleva animal apadrinado.")
    if kind in {"volunteer", "staff"}:
        person = Person.objects.filter(document_number=row["rut"]).first()
        if person and hasattr(person, "team_member"):
            raise CellError("Esta persona ya es parte del equipo o postuló.")


def _check_people_file(rows):
    errors = defaultdict(list)
    team_rows = defaultdict(list)
    for number, row in rows:
        if row["type"] in {"volunteer", "staff"}:
            team_rows[row["rut"]].append(number)
    for numbers in team_rows.values():
        for number in numbers[1:]:
            errors[number].append(f"Esta persona ya aparece como parte del equipo en la fila {numbers[0]}.")
    emails = {}
    for number, row in rows:
        if row["email"]:
            first = emails.setdefault(row["email"], (number, row["rut"]))
            if first[1] != row["rut"]:
                errors[number].append(f"El correo ya aparece con otro RUT en la fila {first[0]}.")
    return errors


def _check_restriction_row(row):
    person = Person.objects.filter(document_number=row["rut"]).first()
    if person and Restriction.objects.filter(person=person, species=row["species"], lifted_at__isnull=True).exists():
        raise CellError("Esta persona ya tiene una restricción vigente para esa especie.")


def _check_restriction_file(rows):
    errors = defaultdict(list)
    seen = {}
    for number, row in rows:
        key = (row["rut"], row["species"].id if row["species"] else None)
        if key in seen:
            errors[number].append(f"Repite la restricción de la fila {seen[key]}.")
        seen.setdefault(key, number)
    return errors


# ---------- Saving ----------

def _person(row, people_cache):
    person = people_cache.get(row["rut"]) or Person.objects.filter(document_number=row["rut"]).first()
    if not person:
        person = Person.objects.create(
            document_number=row["rut"], first_names=row["first_names"], paternal_surname=row["paternal_surname"],
            maternal_surname=row["maternal_surname"], phone=row.get("phone", ""), email=row.get("email") or None,
        )
    people_cache[row["rut"]] = person
    return person


def _save_animals(rows, user):
    for _, row in rows:
        Animal.objects.create(
            slug=unique_slug(row["name"]), name=row["name"], species=row["species"], sex=row["sex"], size=row["size"],
            birth_date_approx=row["birth_date_approx"], intake_date=row["intake_date"],
        )


def _save_adoptions(rows, user):
    people, animals = {}, {}
    for _, row in sorted(rows, key=lambda item: item[1]["contract_signed_at"]):
        key = _animal_key(row)
        if key not in animals:
            animals[key] = Animal.objects.create(
                slug=unique_slug(row["animal_name"]), name=row["animal_name"], species=row["species"],
                birth_date_approx=row["birth_date_approx"], intake_date=row["intake_date"],
            )
        Adoption.objects.create(
            animal=animals[key], person=_person(row, people), source=Adoption.Source.IMPORT,
            status=Adoption.Status.RETURNED if row["returned_at"] else Adoption.Status.SIGNED,
            contract_signed_at=row["contract_signed_at"], fee_charged=row["fee_charged"], discount=0,
            returned_at=row["returned_at"], return_reason=row["return_reason"], reviewed_by=user,
        )


def _save_people(rows, user):
    people = {}
    for _, row in rows:
        person = _person(row, people)
        if row["type"] in {"donation", "sponsorship"}:
            Membership.objects.create(
                person=person, kind=row["type"], monthly_amount=row["monthly_amount"], animal=row.get("animal"),
                status=ReviewStatus.APPROVED, started_at=row["start_date"], reviewed_by=user,
            )
        else:
            TeamMember.objects.create(
                person=person, kind=row["type"], status=ReviewStatus.APPROVED, joined_at=row["start_date"], reviewed_by=user,
            )


def _save_restrictions(rows, user):
    people = {}
    for _, row in rows:
        Restriction.objects.create(
            person=_person(row, people), species=row["species"], reason=row["reason"], notes=row["notes"], created_by=user
        )


KINDS = {
    ImportBatch.Kind.ANIMALS: Kind(
        label="animales",
        columns=[
            Column("Nombre", "name", text), Column("Especie", "species", species_cell),
            Column("Sexo", "sex", choice(SEX), hint="Hembra, Macho o Desconocido"),
            Column("Tamaño", "size", choice(SIZE), hint="Pequeño, Mediano o Grande"),
            *ANIMAL_DATES,
        ],
        row_check=_check_animal_dates, save=_save_animals,
        notes=["Los animales quedan disponibles y aparecen en la web al importarlos."],
    ),
    ImportBatch.Kind.ADOPTIONS: Kind(
        label="adopciones",
        columns=[
            *PERSON_COLUMNS, Column("Teléfono", "phone", phone_cell),
            Column("Animal", "animal_name", text), Column("Especie", "species", species_cell), *ANIMAL_DATES,
            Column("Fecha del contrato", "contract_signed_at", date_cell, hint="Día-mes-año"),
            Column("Aporte cobrado", "fee_charged", amount_cell, hint="En pesos, 0 si no se cobró"),
            Column("Fecha de devolución", "returned_at", optional_date, required=False, hint="Solo si el animal fue devuelto"),
            Column("Motivo de devolución", "return_reason", optional_text, required=False),
        ],
        row_check=_check_adoption_row, file_check=_check_adoption_file, save=_save_adoptions,
        notes=[
            "Varias filas con el mismo animal (nombre, especie, nacimiento y llegada) forman su historial.",
            "Si la última adopción de un animal fue devuelta, el animal vuelve a quedar a cargo de la fundación y aparece en la web.",
        ],
    ),
    ImportBatch.Kind.PEOPLE: Kind(
        label="socios y equipo",
        columns=[
            *PERSON_COLUMNS, Column("Teléfono", "phone", phone_cell), Column("Correo", "email", email_cell, required=False),
            Column("Tipo", "type", choice(PERSON_TYPE), hint="Socio, Padrino, Voluntario o Planta"),
            Column("Monto mensual", "monthly_amount", optional_amount, required=False, hint="Solo socios y padrinos, en pesos"),
            Column("Animal apadrinado", "sponsored_animal", optional_text, required=False, hint="Solo padrinos"),
            Column("Fecha de inicio", "start_date", date_cell, hint="Día-mes-año"),
        ],
        row_check=_check_person_row, file_check=_check_people_file, save=_save_people,
        notes=["Una fila por cada rol: si alguien es socio y voluntario, va en dos filas con el mismo RUT."],
    ),
    ImportBatch.Kind.RESTRICTIONS: Kind(
        label="restricciones",
        columns=[
            *PERSON_COLUMNS, Column("Motivo", "reason", choice(REASON), hint=", ".join(REASON)),
            Column("Especie", "species", lambda v: species_cell(v, allow_all=True), hint="Nombre de la especie o Todas"),
            Column("Notas", "notes", text),
        ],
        row_check=_check_restriction_row, file_check=_check_restriction_file, save=_save_restrictions,
    ),
}


# ---------- Template ----------

def template(kind):
    spec = KINDS[kind]
    book = Workbook()
    sheet = book.active
    sheet.title = "Datos"
    sheet.append([c.header for c in spec.columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, column in enumerate(spec.columns, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(16, len(column.header) + 4)
    help_sheet = book.create_sheet("Instrucciones")
    help_sheet.append(["Columna", "¿Obligatoria?", "Cómo completarla"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for column in spec.columns:
        help_sheet.append([column.header, "Sí" if column.required else "No", column.hint])
    help_sheet.append([])
    help_sheet.append(["Completa la hoja “Datos” sin cambiar los títulos de las columnas."])
    help_sheet.append(["Si una sola fila tiene errores, no se importa nada: la vista previa te dirá qué corregir."])
    for note in spec.notes:
        help_sheet.append([note])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ---------- Validation and import ----------

def validate(kind, uploaded):
    spec = KINDS[kind]
    expected = [c.header for c in spec.columns]
    try:
        book = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception:
        return {"valid": False, "detail": "No pudimos leer el archivo. Sube una planilla Excel (.xlsx) creada desde la plantilla.",
                "missing_columns": [], "extra_columns": [], "rows": []}
    sheet = book["Datos"] if "Datos" in book.sheetnames else book.worksheets[0]
    lines = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in (lines[0] if lines else [])]
    while header and not header[-1]:
        header.pop()
    if header != expected:
        return {
            "valid": False,
            "detail": f"Esta planilla no corresponde a la plantilla de {spec.label}. No se importará nada.",
            "missing_columns": [h for h in expected if h not in header],
            "extra_columns": [h for h in header if h and h not in expected],
            "rows": [],
        }
    parsed, report = [], []
    for number, values in enumerate(lines[1:], start=2):
        values = list(values[: len(expected)]) + [None] * (len(expected) - len(values))
        if all(v in (None, "") for v in values):
            continue
        row, errors = {}, []
        for column, value in zip(spec.columns, values):
            try:
                row[column.key] = column.parse(value)
            except CellError as exc:
                errors.append(f"{column.header}: {exc}")
        if not errors and spec.row_check:
            try:
                spec.row_check(row)
            except CellError as exc:
                errors.append(str(exc))
        report.append({"row": number, "values": [_display(v) for v in values], "errors": errors})
        if not errors:
            parsed.append((number, row))
    if not report:
        return {"valid": False, "detail": "La planilla no tiene filas con datos.", "missing_columns": [], "extra_columns": [], "rows": []}
    if spec.file_check and parsed:
        for number, messages in spec.file_check(parsed).items():
            next(r for r in report if r["row"] == number)["errors"].extend(messages)
    bad = sum(1 for r in report if r["errors"])
    return {
        "valid": bad == 0,
        "detail": "Todas las filas están correctas." if bad == 0 else
        f"{bad} {'fila tiene' if bad == 1 else 'filas tienen'} errores. No se importará nada hasta que todas estén correctas.",
        "missing_columns": [], "extra_columns": [], "rows": report, "_parsed": parsed,
    }


def run(kind, uploaded, user):
    result = validate(kind, uploaded)
    parsed = result.pop("_parsed", None)
    if not result["valid"]:
        return None, result
    with transaction.atomic():
        KINDS[kind].save(parsed, user)
        batch = ImportBatch.objects.create(kind=kind, file_name=uploaded.name[:200], row_count=len(parsed), created_by=user)
    return batch, result


def _display(value):
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    return "" if value is None else str(value)
