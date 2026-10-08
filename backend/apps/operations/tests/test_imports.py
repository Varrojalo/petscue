from datetime import date
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from apps.adoptions.models import Adoption
from apps.animals.models import Animal
from apps.animals.services import in_care
from apps.core.testing import make_animal, next_rut
from apps.core.testing_api import signed_in
from apps.operations import imports
from apps.operations.models import ImportBatch
from apps.people.models import Membership, Person, TeamMember
from apps.restrictions.models import Restriction


def workbook(header, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Datos"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("planilla.xlsx", buffer.getvalue())


def headers(kind):
    return [c.header for c in imports.KINDS[kind].columns]


class ImportApiTests(TestCase):
    def setUp(self):
        self.admin, self.user = signed_in("admin")

    def post(self, kind, step, upload):
        return self.admin.post(f"/api/dashboard/imports/{kind}/{step}", {"file": upload})

    def test_template_has_the_columns_and_instructions(self):
        response = self.admin.get("/api/dashboard/imports/adoptions/template")
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Datos", "Instrucciones"])
        self.assertEqual([c.value for c in book["Datos"][1]], headers("adoptions"))

    def test_only_administration(self):
        worker, _ = signed_in("worker")
        self.assertEqual(worker.get("/api/dashboard/imports/animals/template").status_code, 403)

    def test_wrong_template_imports_nothing(self):
        upload = workbook(["Nombre", "Teléfono", "Dirección"], [["Ana", "123", "Calle 1"]])
        result = self.post("animals", "preview", upload).json()
        self.assertFalse(result["valid"])
        self.assertIn("Fecha de llegada", result["missing_columns"])
        self.assertEqual(result["extra_columns"], ["Teléfono", "Dirección"])

    def test_one_bad_row_imports_nothing(self):
        rows = [
            ["Miel", "Perro", "Hembra", "Mediano", "12-02-2024", "12-02-2026"],
            ["Nieve", "Conejo", "Hembra", "Pequeño", "01-01-2025", "20-08-2026"],
            ["Tito", "Perro", "Macho", "Grande", "01-01-2020", ""],
        ]
        before = Animal.objects.count()
        response = self.post("animals", "confirm", workbook(headers("animals"), rows))
        self.assertEqual(response.status_code, 400)
        errors = {r["row"]: r["errors"] for r in response.json()["rows"]}
        self.assertEqual(errors[2], [])
        self.assertIn("La especie “Conejo” no existe", errors[3][0])
        self.assertEqual(errors[4], ["Fecha de llegada: Falta la fecha."])
        self.assertEqual(Animal.objects.count(), before)
        self.assertFalse(ImportBatch.objects.exists())

    def test_valid_animals_are_imported_and_recorded(self):
        rows = [["Miel", "Perro", "Hembra", "Mediano", "12-02-2024", "12-02-2026"], ["Pancho", "gatos", "Macho", "Pequeño", date(2025, 5, 3), date(2026, 5, 3)]]
        response = self.post("animals", "confirm", workbook(headers("animals"), rows))
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["row_count"], 2)
        self.assertEqual(set(in_care().values_list("name", flat=True)), {"Miel", "Pancho"})

    def test_adoption_history_and_existing_people(self):
        existing = Person.objects.create(document_number=next_rut(), first_names="Rosa", paternal_surname="Díaz")
        r1, r2 = existing.document_number, next_rut()
        rows = [
            [r1, "Otra", "Persona", "", "+56 9 1111 2222", "Chispa", "Perro", "01-01-2020", "01-06-2021", "15-03-2022", 30000, "10-09-2022", "Alergia"],
            [r2, "Luis", "Araya", "", "+56 9 3333 4444", "Chispa", "Perro", "01-01-2020", "01-06-2021", "01-12-2022", 30000, "", ""],
        ]
        response = self.post("adoptions", "confirm", workbook(headers("adoptions"), rows))
        self.assertEqual(response.status_code, 201, response.json())
        chispa = Animal.objects.get(name="Chispa")
        self.assertEqual(list(chispa.adoptions.order_by("contract_signed_at").values_list("status", flat=True)), ["returned", "signed"])
        self.assertEqual(Person.objects.get(pk=existing.pk).first_names, "Rosa")
        self.assertEqual(Adoption.objects.filter(source="import").count(), 2)

    def test_adoption_history_must_be_consistent(self):
        rut_value = next_rut()
        rows = [
            [rut_value, "Ana", "Pérez", "", "+56 9 1111 2222", "Coco", "Perro", "01-01-2020", "01-06-2021", "15-03-2022", 0, "", ""],
            [rut_value, "Ana", "Pérez", "", "+56 9 1111 2222", "Coco", "Perro", "01-01-2020", "01-06-2021", "01-12-2022", 0, "", ""],
        ]
        result = self.post("adoptions", "preview", workbook(headers("adoptions"), rows)).json()
        self.assertFalse(result["valid"])
        self.assertIn("debe tener fecha de devolución", result["rows"][0]["errors"][0])

    def test_people_roles(self):
        animal = make_animal(name="Bruno")
        rut_value = next_rut()
        rows = [
            [rut_value, "Irene", "Salas", "", "+56 9 5555 6666", "irene@example.org", "Socio", 5000, "", "01-03-2025"],
            [rut_value, "Irene", "Salas", "", "+56 9 5555 6666", "irene@example.org", "Voluntario", "", "", "01-03-2025"],
            [next_rut(), "Tomás", "Paredes", "", "+56 9 7777 8888", "", "Padrino", 7000, "bruno", "01-02-2026"],
        ]
        response = self.post("people", "confirm", workbook(headers("people"), rows))
        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(Membership.objects.filter(status="approved").count(), 2)
        self.assertEqual(Membership.objects.get(kind="sponsorship").animal, animal)
        self.assertEqual(TeamMember.objects.get().person.email, "irene@example.org")

    def test_restrictions(self):
        rows = [[next_rut(), "Óscar", "Méndez", "", "Abandono", "Todas", "Abandonó a un perro en 2023"],
                [next_rut(), "Rita", "Fuentes", "", "Maltrato", "Perros", "Denuncia"]]
        response = self.post("restrictions", "confirm", workbook(headers("restrictions"), rows))
        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(Restriction.objects.filter(species__isnull=True).count(), 1)
        self.assertEqual(Restriction.objects.get(person__first_names="Rita").created_by, self.user)
